import os
import io
import re
import html
import json as _json_mod
import socket
import ipaddress
from datetime import datetime, timedelta
from urllib.parse import urlparse

import httpx
from fastapi import FastAPI, HTTPException, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import PlainTextResponse, HTMLResponse, Response, JSONResponse
from pydantic import BaseModel, field_validator, Field


app = FastAPI(title="10blackfingers API", version="2.1.0")

ALLOWED_ORIGINS = [
    "https://10blackfingers.vercel.app",
    "http://localhost:3000",
    "http://127.0.0.1:3000",
    "http://localhost:5500",
    "http://127.0.0.1:5500",
    "http://localhost:8080",
    "http://127.0.0.1:8080",
]

app.add_middleware(
    CORSMiddleware,
    allow_origins=ALLOWED_ORIGINS,
    allow_credentials=False,
    allow_methods=["GET", "POST"],
    allow_headers=["Content-Type"],
)


SUPABASE_URL = os.getenv("SUPABASE_URL", "")
SUPABASE_SERVICE_KEY = os.getenv("SUPABASE_SERVICE_KEY", "")

DISABLE_SCAN_LIMITS = os.getenv("DISABLE_SCAN_LIMITS", "0") == "1"

MAX_SCANS_PER_HOUR = 5
MAX_URL_LENGTH = 2048
DNS_TIMEOUT_SECONDS = 5
RESULT_DEFAULT_LIMIT = 500
RESULT_MAX_LIMIT = 2000

DANGEROUS_URL_CHARS = set("'`;|<>\\\n\r")


def sb_headers():
    return {
        "apikey": SUPABASE_SERVICE_KEY,
        "Authorization": f"Bearer {SUPABASE_SERVICE_KEY}",
        "Content-Type": "application/json",
        "Prefer": "return=representation",
    }


async def sb_insert(table: str, data: dict):
    url = f"{SUPABASE_URL}/rest/v1/{table}"
    async with httpx.AsyncClient(timeout=30.0) as client:
        r = await client.post(url, headers=sb_headers(), json=data)
        if r.status_code >= 400:
            raise Exception(f"Supabase insert error: {r.status_code} - {r.text}")
        return r.json()


async def sb_select(table: str, params: dict = None):
    url = f"{SUPABASE_URL}/rest/v1/{table}"
    async with httpx.AsyncClient(timeout=30.0) as client:
        r = await client.get(url, headers=sb_headers(), params=params or {})
        if r.status_code >= 400:
            raise Exception(f"Supabase select error: {r.status_code} - {r.text}")
        return r.json()


async def sb_update(table: str, match: dict, data: dict):
    params = {k: f"eq.{v}" for k, v in match.items()}
    url = f"{SUPABASE_URL}/rest/v1/{table}"
    async with httpx.AsyncClient(timeout=30.0) as client:
        r = await client.patch(url, headers=sb_headers(), params=params, json=data)
        if r.status_code >= 400:
            raise Exception(f"Supabase update error: {r.status_code} - {r.text}")
        return r.json()


BLOCKED_HOSTNAMES = {"localhost", "metadata.google.internal"}


def is_private_or_reserved_ip(ip_str: str) -> bool:
    try:
        ip = ipaddress.ip_address(ip_str)
    except ValueError:
        return True
    return (
        ip.is_private
        or ip.is_loopback
        or ip.is_link_local
        or ip.is_multicast
        or ip.is_reserved
        or ip.is_unspecified
    )


def validate_target_url(raw_url: str) -> str:
    raw_url = (raw_url or "").strip()

    if not raw_url:
        raise HTTPException(status_code=400, detail="الرابط مطلوب")

    if len(raw_url) > MAX_URL_LENGTH:
        raise HTTPException(status_code=400, detail=f"الرابط طويل جداً (الحد الأقصى {MAX_URL_LENGTH} حرف)")

    if any(c in raw_url for c in DANGEROUS_URL_CHARS):
        raise HTTPException(status_code=400, detail="الرابط يحتوي على أحرف غير مسموحة")

    parsed = urlparse(raw_url)

    if parsed.scheme not in ("http", "https"):
        raise HTTPException(status_code=400, detail="الرابط لازم يبدأ بـ http:// أو https://")

    hostname = parsed.hostname
    if not hostname:
        raise HTTPException(status_code=400, detail="رابط غير صالح")

    hostname_lower = hostname.lower()
    if hostname_lower in BLOCKED_HOSTNAMES or hostname_lower.endswith(".local"):
        raise HTTPException(status_code=400, detail="غير مسموح بفحص عناوين داخلية")

    try:
        ipaddress.ip_address(hostname)
        if is_private_or_reserved_ip(hostname):
            raise HTTPException(status_code=400, detail="غير مسموح بفحص عناوين IP داخلية أو محجوزة")
    except ValueError:
        pass

    original_timeout = socket.getdefaulttimeout()
    try:
        socket.setdefaulttimeout(DNS_TIMEOUT_SECONDS)
        resolved_ips = socket.getaddrinfo(hostname, None, proto=socket.IPPROTO_TCP)
        for family, _, _, _, sockaddr in resolved_ips:
            ip_candidate = sockaddr[0]
            if is_private_or_reserved_ip(ip_candidate):
                raise HTTPException(
                    status_code=400,
                    detail="الدومين بيشاور على عنوان IP داخلي - غير مسموح بالفحص",
                )
    except socket.timeout:
        raise HTTPException(status_code=400, detail="تعذر إيجاد الدومين (DNS timeout)")
    except socket.gaierror:
        raise HTTPException(status_code=400, detail="تعذر إيجاد الدومين (DNS resolution فشل)")
    finally:
        socket.setdefaulttimeout(original_timeout)

    return raw_url


def validate_session_id(session_id: str) -> str:
    session_id = (session_id or "").strip()
    if not session_id or len(session_id) > 100:
        raise HTTPException(status_code=400, detail="session_id غير صالح")
    if not re.match(r"^[A-Za-z0-9_\-]+$", session_id):
        raise HTTPException(status_code=400, detail="session_id يحتوي على أحرف غير مسموحة")
    return session_id


async def check_rate_limit(session_id: str):
    if DISABLE_SCAN_LIMITS:
        return
    try:
        one_hour_ago = (datetime.utcnow() - timedelta(hours=1)).isoformat()
        recent = await sb_select("scans", {
            "session_id": f"eq.{session_id}",
            "created_at": f"gte.{one_hour_ago}",
            "status": "neq.failed",
            "select": "id",
        })
        if len(recent or []) >= MAX_SCANS_PER_HOUR:
            raise HTTPException(
                status_code=429,
                detail=f"تجاوزت الحد المسموح من الفحوصات ({MAX_SCANS_PER_HOUR} في الساعة). حاول لاحقاً",
            )
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(
            status_code=503,
            detail="تعذر التحقق من الحد المسموح، حاول بعد قليل",
        )


async def check_concurrent_scans(session_id: str):
    if DISABLE_SCAN_LIMITS:
        return
    try:
        active_scans = await sb_select("scans", {
            "session_id": f"eq.{session_id}",
            "status": "in.(pending,running)",
            "select": "id,status,created_at",
        })

        if not active_scans:
            return

        now = datetime.utcnow()
        truly_active = []

        for scan in active_scans:
            scan_id = scan.get("id", "")
            status = scan.get("status", "")
            created_at_str = scan.get("created_at", "")

            age_minutes = 0.0
            try:
                if created_at_str.endswith("Z"):
                    created_at_str = created_at_str[:-1] + "+00:00"
                scan_time = datetime.fromisoformat(created_at_str)
                if scan_time.tzinfo:
                    age_minutes = (datetime.now(scan_time.tzinfo) - scan_time).total_seconds() / 60
                else:
                    age_minutes = (now - scan_time).total_seconds() / 60
            except Exception:
                age_minutes = 0.0

            threshold = 5 if status == "pending" else 185

            if age_minutes < threshold:
                truly_active.append(scan)
            else:
                try:
                    await sb_update("scans", {"id": scan_id}, {
                        "status": "failed",
                        "error_message": f"Auto-cleaned: stuck in '{status}' for {int(age_minutes)} min",
                        "completed_at": now.isoformat(),
                    })
                except Exception:
                    pass

        if truly_active:
            raise HTTPException(
                status_code=429,
                detail="لديك فحص قيد التنفيذ بالفعل - انتظر انتهاءه قبل بدء فحص جديد",
            )
    except HTTPException:
        raise
    except Exception:
        pass


def secure_report_headers(scan_id: str, ext: str, attachment: bool = True) -> dict:
    headers = {
        "Cache-Control": "no-store, no-cache, must-revalidate, private",
        "Pragma": "no-cache",
        "X-Content-Type-Options": "nosniff",
    }
    if attachment:
        headers["Content-Disposition"] = f'attachment; filename="report-{scan_id}.{ext}"'
    return headers


class ScanRequest(BaseModel):
    url: str = Field(..., max_length=MAX_URL_LENGTH)
    consent: bool
    session_id: str

    @field_validator("session_id")
    @classmethod
    def validate_session_id_field(cls, v):
        v = (v or "").strip()
        if not v or len(v) > 100 or not re.match(r"^[A-Za-z0-9_\-]+$", v):
            raise ValueError("session_id غير صالح")
        return v


class ScanResponse(BaseModel):
    scan_id: str
    status: str
    message: str


@app.get("/")
def root():
    return {"name": "10blackfingers API", "version": "2.1.0", "status": "running"}


@app.get("/api")
def api_root():
    return {"name": "10blackfingers API", "version": "2.1.0", "status": "running"}


@app.get("/health")
def health_check():
    return {"status": "healthy", "timestamp": datetime.utcnow().isoformat()}


@app.post("/api/scan", response_model=ScanResponse)
async def create_scan(request: ScanRequest):
    import uuid

    if not request.consent:
        raise HTTPException(status_code=400, detail="الموافقة على التصريح مطلوبة قبل الفحص")

    await check_rate_limit(request.session_id)
    await check_concurrent_scans(request.session_id)

    target_url = validate_target_url(request.url)

    scan_id = str(uuid.uuid4())

    try:
        await sb_insert("scans", {
            "id": scan_id,
            "target_url": target_url,
            "status": "pending",
            "session_id": request.session_id,
            "current_phase": "queued",
            "progress_percent": 0,
            "elapsed_seconds": 0,
        })
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"DB error: {str(e)}")

    github_token = os.getenv("GITHUB_TOKEN", "")
    github_repo = os.getenv("GITHUB_REPO", "")

    if not github_token:
        try:
            await sb_update("scans", {"id": scan_id}, {
                "status": "failed",
                "error_message": "GITHUB_TOKEN غير مُعرّف في إعدادات السيرفر",
            })
        except Exception:
            pass
        raise HTTPException(
            status_code=500,
            detail="GITHUB_TOKEN غير مُعرّف في إعدادات السيرفر - تواصل مع المسؤول",
        )

    if not github_repo or "/" not in github_repo or len(github_repo.split("/")) != 2:
        try:
            await sb_update("scans", {"id": scan_id}, {
                "status": "failed",
                "error_message": "GITHUB_REPO غير مُعرّف أو غير صالح في إعدادات السيرفر",
            })
        except Exception:
            pass
        raise HTTPException(
            status_code=500,
            detail="GITHUB_REPO غير مُعرّف أو غير صالح - يجب أن يكون بصيغة owner/repo",
        )

    try:
        url = f"https://api.github.com/repos/{github_repo}/actions/workflows/security-scan.yml/dispatches"
        headers = {
            "Authorization": f"Bearer {github_token}",
            "Accept": "application/vnd.github+json",
            "X-GitHub-Api-Version": "2022-11-28",
        }
        payload = {"ref": "main", "inputs": {"scan_id": scan_id, "target_url": target_url}}
        async with httpx.AsyncClient(timeout=30.0) as client:
            dispatch_res = await client.post(url, headers=headers, json=payload)

        if dispatch_res.status_code >= 300:
            log_detail = f"GitHub dispatch failed: {dispatch_res.status_code} - {dispatch_res.text[:300]}"
            try:
                await sb_update("scans", {"id": scan_id}, {
                    "status": "failed",
                    "error_message": log_detail,
                })
            except Exception:
                pass
            raise HTTPException(
                status_code=502,
                detail=f"تعذر بدء الفحص على GitHub (HTTP {dispatch_res.status_code}). تحقق من إعدادات GITHUB_TOKEN و GITHUB_REPO",
            )
    except HTTPException:
        raise
    except Exception as e:
        try:
            await sb_update("scans", {"id": scan_id}, {
                "status": "failed",
                "error_message": f"تعذر بدء الفحص: {str(e)}",
            })
        except Exception:
            pass
        raise HTTPException(status_code=500, detail=f"تعذر بدء الفحص: {str(e)}")

    return ScanResponse(scan_id=scan_id, status="pending", message="تم إضافة الفحص للطابور بنجاح")


@app.get("/api/scan/{scan_id}")
async def get_scan_status(
    scan_id: str,
    limit: int = Query(RESULT_DEFAULT_LIMIT, ge=1, le=RESULT_MAX_LIMIT),
    session_id: str = Query(...),
):
    try:
        session_id = validate_session_id(session_id)

        scans = await sb_select("scans", {
            "id": f"eq.{scan_id}",
            "session_id": f"eq.{session_id}",
            "select": "*",
        })
        if not scans:
            raise HTTPException(status_code=404, detail="الفحص غير موجود أو لا تملك صلاحية الوصول")

        scan = scans[0]
        results = await sb_select("results", {
            "scan_id": f"eq.{scan_id}",
            "select": "*",
            "limit": str(limit),
        })
        return {"scan": scan, "results": results or []}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Fetch error: {str(e)}")


@app.get("/api/scan/{scan_id}/status")
async def get_scan_live_status(scan_id: str):
    try:
        scans = await sb_select("scans", {
            "id": f"eq.{scan_id}",
            "select": "id,status,current_phase,progress_percent,elapsed_seconds,total_findings,critical_count,high_count,medium_count,low_count,started_at,completed_at,error_message,target_url",
        })
        if not scans:
            raise HTTPException(status_code=404, detail="الفحص غير موجود")
        return {"scan": scans[0]}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Status error: {str(e)}")


@app.get("/api/scan/{scan_id}/results")
async def get_scan_results(
    scan_id: str,
    offset: int = Query(0, ge=0),
    limit: int = Query(RESULT_DEFAULT_LIMIT, ge=1, le=RESULT_MAX_LIMIT),
    severity: str = Query(None),
    session_id: str = Query(...),
):
    try:
        session_id = validate_session_id(session_id)

        scan_check = await sb_select("scans", {
            "id": f"eq.{scan_id}",
            "session_id": f"eq.{session_id}",
            "select": "id",
        })
        if not scan_check:
            raise HTTPException(status_code=404, detail="الفحص غير موجود أو لا تملك صلاحية الوصول")

        params = {
            "scan_id": f"eq.{scan_id}",
            "select": "*",
            "offset": str(offset),
            "limit": str(limit),
            "order": "severity.asc,created_at.asc",
        }
        if severity and severity in ("critical", "high", "medium", "low", "info"):
            params["severity"] = f"eq.{severity}"

        results = await sb_select("results", params)

        count_params = {"scan_id": f"eq.{scan_id}", "select": "id"}
        if severity:
            count_params["severity"] = f"eq.{severity}"
        all_results = await sb_select("results", count_params)

        total = len(all_results or [])
        return {
            "results": results or [],
            "offset": offset,
            "limit": limit,
            "total": total,
            "has_more": offset + len(results or []) < total,
        }
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Results error: {str(e)}")


@app.get("/api/scans")
async def list_scans(session_id: str, limit: int = 10):
    session_id = validate_session_id(session_id)
    limit = max(1, min(limit, 50))
    try:
        scans = await sb_select("scans", {
            "session_id": f"eq.{session_id}",
            "select": "*",
            "order": "created_at.desc",
            "limit": str(limit),
        })
        return {"scans": scans or []}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"List error: {str(e)}")


async def fetch_report_data(scan_id: str, session_id: str):
    session_id = validate_session_id(session_id)
    scans = await sb_select("scans", {
        "id": f"eq.{scan_id}",
        "session_id": f"eq.{session_id}",
        "select": "*",
    })
    if not scans:
        raise HTTPException(status_code=404, detail="الفحص غير موجود أو لا تملك صلاحية الوصول")
    scan = scans[0]
    results = await sb_select("results", {"scan_id": f"eq.{scan_id}", "select": "*"})
    return scan, results or []


def _esc(value) -> str:
    if value is None:
        return ""
    return html.escape(str(value), quote=True)


def _fmt_val(value, max_len=3000):
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        try:
            s = _json_mod.dumps(value, ensure_ascii=False, default=str)
        except Exception:
            s = str(value)
    else:
        s = str(value)
    if len(s) > max_len:
        s = s[:max_len] + "... (مقتطع)"
    return s


SEVERITY_ICONS = {
    "critical": "🔴",
    "high": "🟠",
    "medium": "🟡",
    "low": "🟢",
    "info": "ℹ️",
}

SEVERITY_AR = {
    "critical": "ثغرات حرجة",
    "high": "ثغرات عالية",
    "medium": "ثغرات متوسطة",
    "low": "ثغرات منخفضة",
    "info": "معلومات",
}


def build_txt_report(scan, results):
    lines = []
    lines.append("=" * 78)
    lines.append("    10blackfingers - Security Assessment Report")
    lines.append("=" * 78)
    lines.append(f"Target URL:     {scan.get('target_url', 'N/A')}")
    lines.append(f"Scan ID:        {scan.get('id', 'N/A')}")
    lines.append(f"Started:        {scan.get('started_at', 'N/A')}")
    lines.append(f"Completed:      {scan.get('completed_at', 'N/A')}")
    lines.append(f"Duration:       {scan.get('elapsed_seconds', 'N/A')} seconds")
    lines.append(f"Status:         {scan.get('status', 'N/A')}")
    lines.append("")
    lines.append("-" * 78)
    lines.append("  EXECUTIVE SUMMARY")
    lines.append("-" * 78)
    lines.append(f"Total Findings:  {scan.get('total_findings', 0)}")
    lines.append(f"  Critical:      {scan.get('critical_count', 0)}")
    lines.append(f"  High:          {scan.get('high_count', 0)}")
    lines.append(f"  Medium:        {scan.get('medium_count', 0)}")
    lines.append(f"  Low:           {scan.get('low_count', 0)}")
    lines.append("")

    by_severity = {"critical": [], "high": [], "medium": [], "low": [], "info": []}
    for r in results:
        by_severity.setdefault(r.get("severity", "info"), []).append(r)

    for severity in ["critical", "high", "medium", "low", "info"]:
        items = by_severity.get(severity, [])
        if not items:
            continue
        icon = SEVERITY_ICONS.get(severity, "")
        lines.append("")
        lines.append("=" * 78)
        lines.append(f"  {icon} {severity.upper()} FINDINGS ({len(items)})")
        lines.append("=" * 78)

        for i, r in enumerate(items, 1):
            lines.append("")
            lines.append(f"[{i}] {r.get('title', 'Untitled')}")
            lines.append("-" * 60)
            lines.append(f"  Tool:       {r.get('tool_name', 'N/A')}")
            lines.append(f"  Phase:      {r.get('phase', 'N/A')}")
            lines.append(f"  Severity:   {r.get('severity', 'N/A')}")
            if r.get("cvss_score"):
                lines.append(f"  CVSS Score: {r['cvss_score']}")
            if r.get("cve_id"):
                lines.append(f"  CVE:        {r['cve_id']}")
            if r.get("description"):
                lines.append("")
                lines.append("  DESCRIPTION:")
                for line in str(r["description"]).split("\n"):
                    lines.append(f"    {line}")

            ev = r.get("evidence")
            if ev:
                lines.append("")
                lines.append("  EVIDENCE & EXTRACTED DATA:")
                if isinstance(ev, dict):
                    for k, v in ev.items():
                        if k in ("extracted_pii", "aggregated") and isinstance(v, dict):
                            lines.append(f"    [{k}]")
                            for pk, pv in v.items():
                                if isinstance(pv, list):
                                    for pvv in pv:
                                        lines.append(f"      {pk}: {pvv}")
                                else:
                                    lines.append(f"      {pk}: {pv}")
                        elif k == "pii_data" and isinstance(v, list):
                            for item in v[:10]:
                                if isinstance(item, dict):
                                    lines.append(f"    [PII from {item.get('source_url', 'N/A')}]")
                                    for pk, pv in item.items():
                                        if pk == "source_url":
                                            continue
                                        lines.append(f"      {pk}: {pv}")
                        elif k == "dumped_files_preview" and isinstance(v, dict):
                            lines.append(f"    [DB DUMP FILES: {len(v)}]")
                            for fn, content in list(v.items())[:10]:
                                lines.append(f"      --- {fn} ---")
                                for line in str(content)[:3000].split("\n"):
                                    lines.append(f"      {line}")
                        elif k == "extracted_files" and isinstance(v, list):
                            lines.append(f"    [EXTRACTED FILES: {len(v)}]")
                            for f in v[:10]:
                                if isinstance(f, dict):
                                    lines.append(f"      - {f.get('file', '')}: {f.get('url', '')}")
                                    for line in str(f.get("content", ""))[:2000].split("\n"):
                                        lines.append(f"        {line}")
                        else:
                            lines.append(f"    {k}: {_fmt_val(v)}")
                else:
                    lines.append(f"    {str(ev)[:3000]}")

            if r.get("remediation"):
                lines.append("")
                lines.append("  REMEDIATION:")
                for line in str(r["remediation"]).split("\n"):
                    lines.append(f"    {line}")

    lines.append("")
    lines.append("=" * 78)
    lines.append("  END OF REPORT")
    lines.append("=" * 78)
    lines.append(f"Generated by 10blackfingers at {datetime.utcnow().isoformat()} UTC")
    return "\n".join(lines)


def _render_pii_table_html(title, pii_dict):
    if not pii_dict or not isinstance(pii_dict, dict):
        return ""
    rows = []
    for key, value in pii_dict.items():
        if key == "source_url":
            continue
        if isinstance(value, list):
            for v in value:
                rows.append((key, str(v)))
        elif isinstance(value, (int, float, str)):
            rows.append((key, str(value)))
        else:
            rows.append((key, _fmt_val(value)))
    if not rows:
        return ""
    html_out = f'<div class="pii-block"><h4 class="pii-title">🔐 {_esc(title)} ({len(rows)})</h4>'
    html_out += '<table class="data-table"><thead><tr><th style="width:30%">النوع</th><th>القيمة</th></tr></thead><tbody>'
    for k, v in rows[:500]:
        html_out += f'<tr><td><strong>{_esc(k)}</strong></td><td><code>{_esc(v)}</code></td></tr>'
    html_out += '</tbody></table>'
    if len(rows) > 500:
        html_out += f'<p class="more-note">... و {len(rows) - 500} صف إضافي في التقرير الكامل</p>'
    html_out += '</div>'
    return html_out


def _render_evidence_html(evidence):
    if not evidence:
        return ""
    if isinstance(evidence, str):
        return f'<pre class="raw-evidence">{_esc(evidence)}</pre>'
    if not isinstance(evidence, dict):
        return ""

    parts = ['<div class="evidence-block"><h4 class="evidence-title">🔍 الأدلة والبيانات المستخرجة</h4>']

    for key in ("extracted_pii", "aggregated"):
        if key in evidence and evidence[key]:
            title = "البيانات الحساسة المستخرجة (PII)" if key == "extracted_pii" else "بيانات PII المجمّعة"
            parts.append(_render_pii_table_html(title, evidence[key]))

    if "pii_data" in evidence and isinstance(evidence["pii_data"], list):
        for i, item in enumerate(evidence["pii_data"][:10], 1):
            if isinstance(item, dict):
                source = item.get("source_url", f"مصدر {i}")
                pii_only = {k: v for k, v in item.items() if k != "source_url"}
                parts.append(_render_pii_table_html(f"بيانات PII من {source}", pii_only))

    if "dumped_files_preview" in evidence and isinstance(evidence["dumped_files_preview"], dict):
        parts.append('<div class="db-dump-block"><h4 class="db-title">🗄️ ملفات قاعدة البيانات المسحوبة</h4>')
        for filename, content in list(evidence["dumped_files_preview"].items())[:20]:
            parts.append('<div class="db-file">')
            parts.append(f'<h5 class="db-file-name">📄 {_esc(filename)}</h5>')
            parts.append(f'<pre class="db-content">{_esc(str(content)[:8000])}</pre>')
            parts.append('</div>')
        parts.append('</div>')

    if "extracted_files" in evidence and isinstance(evidence["extracted_files"], list):
        parts.append('<div class="files-block"><h4 class="files-title">📁 الملفات الحساسة المسحوبة</h4>')
        for f in evidence["extracted_files"][:10]:
            if isinstance(f, dict):
                parts.append('<div class="file-item">')
                parts.append(f'<h5>{_esc(f.get("file", ""))} — <code>{_esc(f.get("url", ""))}</code></h5>')
                if f.get("extracted_pii"):
                    parts.append(_render_pii_table_html("PII داخل الملف", f["extracted_pii"]))
                parts.append(f'<pre>{_esc(str(f.get("content", ""))[:5000])}</pre>')
                parts.append('</div>')
        parts.append('</div>')

    skip_keys = {"extracted_pii", "aggregated", "pii_data", "dumped_files_preview", "extracted_files"}
    generic = {k: v for k, v in evidence.items() if k not in skip_keys and v not in (None, "", [], {})}
    if generic:
        parts.append('<ul class="generic-evidence">')
        for k, v in generic.items():
            parts.append(f'<li><strong>{_esc(k)}:</strong> <code>{_esc(_fmt_val(v))}</code></li>')
        parts.append('</ul>')

    parts.append('</div>')
    return "".join(parts)


def build_html_report(scan, results):
    by_severity = {"critical": [], "high": [], "medium": [], "low": [], "info": []}
    for r in results:
        by_severity.setdefault(r.get("severity", "info"), []).append(r)

    findings_html = ""
    for severity in ["critical", "high", "medium", "low", "info"]:
        items = by_severity.get(severity, [])
        if not items:
            continue
        icon = SEVERITY_ICONS.get(severity, "")
        label_ar = SEVERITY_AR.get(severity, severity.upper())

        findings_html += f'<h2 class="severity-{severity}">{icon} {_esc(label_ar)} ({len(items)})</h2>'
        for i, r in enumerate(items, 1):
            evidence_html = _render_evidence_html(r.get("evidence"))
            cvss_span = f'<span><strong>CVSS:</strong> {_esc(r["cvss_score"])}</span>' if r.get("cvss_score") else ""
            cve_span = f'<span><strong>CVE:</strong> {_esc(r["cve_id"])}</span>' if r.get("cve_id") else ""
            desc_p = f'<p class="desc">{_esc(r["description"])}</p>' if r.get("description") else ""
            rem_div = ""
            if r.get("remediation"):
                rem_div = f'<div class="remediation"><h4>🛠️ كيفية الإصلاح</h4><pre>{_esc(r["remediation"])}</pre></div>'

            findings_html += f"""
<div class="finding severity-{severity}">
    <h3>[{i}] {_esc(r.get('title', 'Untitled'))}</h3>
    <div class="meta">
        <span><strong>الأداة:</strong> {_esc(r.get('tool_name', 'N/A'))}</span>
        <span><strong>المرحلة:</strong> {_esc(r.get('phase', 'N/A'))}</span>
        <span><strong>الخطورة:</strong> {_esc(r.get('severity', 'N/A'))}</span>
        {cvss_span}{cve_span}
    </div>
    {desc_p}
    {evidence_html}
    {rem_div}
</div>
"""

    total = scan.get("total_findings", 0)
    critical = scan.get("critical_count", 0)
    high = scan.get("high_count", 0)
    medium = scan.get("medium_count", 0)
    low = scan.get("low_count", 0)

    if critical > 0:
        overall_msg, overall_class = "خطر — يوجد ثغرات حرجة تحتاج إصلاح فوري", "danger"
    elif high > 0:
        overall_msg, overall_class = "يحتاج انتباه — يوجد ثغرات عالية", "warning"
    elif medium > 0:
        overall_msg, overall_class = "متوسط — يوجد ثغرات متوسطة الخطورة", "medium"
    else:
        overall_msg, overall_class = "آمن — لا يوجد ثغرات خطيرة", "safe"

    target_safe = _esc(scan.get("target_url", "N/A"))

    html_out = f"""<!DOCTYPE html>
<html lang="ar" dir="rtl">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>تقرير أمني — {target_safe}</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Cairo:wght@400;600;700&display=swap" rel="stylesheet">
<style>
* {{ box-sizing: border-box; }}
body {{
    font-family: 'Cairo', -apple-system, Tahoma, sans-serif;
    max-width: 1200px;
    margin: 30px auto;
    padding: 20px;
    background: #f8fafc;
    color: #1f2937;
    direction: rtl;
    line-height: 1.7;
}}
h1 {{
    color: #00a832;
    border-bottom: 3px solid #00a832;
    padding-bottom: 12px;
    font-size: 26px;
}}
h2 {{
    margin-top: 40px;
    padding: 14px 20px;
    border-radius: 8px;
    font-size: 20px;
    font-weight: 700;
}}
h2.severity-critical {{ background: #fee2e2; color: #991b1b; border-right: 5px solid #dc2626; }}
h2.severity-high {{ background: #ffedd5; color: #9a3412; border-right: 5px solid #ea580c; }}
h2.severity-medium {{ background: #fef9c3; color: #854d0e; border-right: 5px solid #ca8a04; }}
h2.severity-low {{ background: #dcfce7; color: #166534; border-right: 5px solid #16a34a; }}
h2.severity-info {{ background: #e0e7ff; color: #3730a3; border-right: 5px solid #6366f1; }}

.overall {{ padding: 18px 24px; border-radius: 10px; font-size: 17px; font-weight: 700; margin: 20px 0; text-align: center; }}
.overall.danger {{ background: #fef2f2; color: #991b1b; border: 2px solid #dc2626; }}
.overall.warning {{ background: #fff7ed; color: #9a3412; border: 2px solid #ea580c; }}
.overall.medium {{ background: #fefce8; color: #854d0e; border: 2px solid #ca8a04; }}
.overall.safe {{ background: #f0fdf4; color: #166534; border: 2px solid #16a34a; }}

.summary {{
    background: #ffffff;
    padding: 24px;
    border-radius: 10px;
    margin: 20px 0;
    box-shadow: 0 2px 10px rgba(0,0,0,0.06);
    border: 1px solid #e5e7eb;
}}
.summary p {{ margin: 8px 0; }}
.stats {{ display: grid; grid-template-columns: repeat(auto-fit, minmax(120px, 1fr)); gap: 12px; margin-top: 20px; }}
.stat {{ padding: 16px; border-radius: 8px; color: white; text-align: center; }}
.stat .num {{ font-size: 24px; font-weight: 700; }}
.stat .lbl {{ font-size: 12px; opacity: 0.95; }}
.stat.total {{ background: #1e293b; }}
.stat.critical {{ background: #dc2626; }}
.stat.high {{ background: #ea580c; }}
.stat.medium {{ background: #ca8a04; }}
.stat.low {{ background: #16a34a; }}

.finding {{
    background: #ffffff;
    padding: 20px 24px;
    border-radius: 10px;
    margin: 16px 0;
    border-right: 5px solid #cbd5e1;
    box-shadow: 0 1px 4px rgba(0,0,0,0.05);
}}
.finding.severity-critical {{ border-right-color: #dc2626; }}
.finding.severity-high {{ border-right-color: #ea580c; }}
.finding.severity-medium {{ border-right-color: #ca8a04; }}
.finding.severity-low {{ border-right-color: #16a34a; }}
.finding.severity-info {{ border-right-color: #6366f1; }}
.finding h3 {{ margin: 0 0 12px 0; color: #111827; font-size: 17px; font-weight: 600; }}

.meta {{
    display: flex; gap: 16px; flex-wrap: wrap;
    font-size: 13px; color: #6b7280;
    margin: 8px 0; padding: 8px 0;
    border-bottom: 1px dashed #e5e7eb;
}}
.desc {{
    background: #f9fafb; padding: 14px; border-radius: 6px;
    margin: 12px 0; color: #374151; font-size: 14px;
}}

.evidence-block {{
    background: #f1f5f9;
    padding: 16px;
    border-radius: 8px;
    margin: 14px 0;
    border: 1px solid #cbd5e1;
}}
.evidence-title {{
    margin: 0 0 12px 0;
    color: #1e40af;
    font-size: 15px;
    border-bottom: 2px solid #3b82f6;
    padding-bottom: 6px;
}}
.pii-block {{
    background: #fff;
    border: 1px solid #fca5a5;
    border-radius: 8px;
    padding: 12px;
    margin: 12px 0;
}}
.pii-title {{ margin: 0 0 10px 0; color: #991b1b; font-size: 14px; }}
.data-table {{ width: 100%; border-collapse: collapse; font-size: 12px; background: #fff; direction: ltr; }}
.data-table thead th {{ background: #1e293b; color: #fff; padding: 8px 10px; text-align: left; font-weight: 600; }}
.data-table tbody td {{ padding: 7px 10px; border-bottom: 1px solid #e5e7eb; word-break: break-all; }}
.data-table tbody tr:nth-child(even) {{ background: #f9fafb; }}
.data-table tbody tr:hover {{ background: #fef3c7; }}
.data-table code {{ color: #b91c1c; font-family: 'Courier New', monospace; font-size: 11px; background: #fef2f2; padding: 2px 5px; border-radius: 3px; }}
.more-note {{ margin: 8px 0 0 0; font-size: 11px; color: #6b7280; font-style: italic; }}

.db-dump-block {{ background: #1e293b; color: #e2e8f0; padding: 14px; border-radius: 8px; margin: 12px 0; }}
.db-title {{ margin: 0 0 12px 0; color: #4ade80; font-size: 15px; }}
.db-file {{ margin: 12px 0; padding: 10px; background: #0f172a; border-radius: 6px; }}
.db-file-name {{ margin: 0 0 8px 0; color: #fbbf24; font-size: 13px; }}
.db-content {{ background: #000; color: #d1fae5; padding: 10px; border-radius: 4px; overflow-x: auto; font-size: 11px; direction: ltr; text-align: left; max-height: 400px; overflow-y: auto; margin: 0; white-space: pre-wrap; word-break: break-all; }}

.files-block {{ background: #fef3c7; border: 1px solid #fbbf24; border-radius: 8px; padding: 12px; margin: 12px 0; }}
.files-title {{ margin: 0 0 10px 0; color: #92400e; font-size: 15px; }}
.file-item {{ background: #fff; padding: 10px; border-radius: 6px; margin: 8px 0; border-right: 3px solid #f59e0b; }}
.file-item h5 {{ margin: 0 0 6px 0; font-size: 12px; color: #78350f; }}
.file-item pre {{ background: #1e293b; color: #e2e8f0; padding: 8px; border-radius: 4px; overflow-x: auto; font-size: 10px; direction: ltr; text-align: left; margin: 6px 0 0 0; max-height: 300px; overflow-y: auto; white-space: pre-wrap; word-break: break-all; }}

.generic-evidence {{ list-style: none; padding: 0; margin: 10px 0; background: #fff; border-radius: 6px; border: 1px solid #e5e7eb; }}
.generic-evidence li {{ padding: 8px 12px; border-bottom: 1px solid #f3f4f6; font-size: 12px; word-break: break-all; }}
.generic-evidence li:last-child {{ border-bottom: none; }}
.generic-evidence code {{ color: #0369a1; background: #f0f9ff; padding: 2px 5px; border-radius: 3px; font-family: 'Courier New', monospace; font-size: 11px; direction: ltr; display: inline-block; }}
.raw-evidence {{ background: #1e293b; color: #e2e8f0; padding: 12px; border-radius: 6px; overflow-x: auto; font-size: 11px; direction: ltr; text-align: left; white-space: pre-wrap; word-break: break-all; }}

.remediation {{ background: #f0fdf4; padding: 14px; border-radius: 8px; margin: 14px 0; border-right: 4px solid #16a34a; }}
.remediation h4 {{ margin: 0 0 10px 0; color: #166534; font-size: 14px; }}
.remediation pre {{ background: #1e293b; color: #e2e8f0; padding: 12px; border-radius: 6px; overflow-x: auto; white-space: pre-wrap; word-break: break-word; margin: 0; font-size: 12px; direction: ltr; text-align: left; }}

.print-btn {{
    position: fixed;
    top: 20px;
    left: 20px;
    background: #00a832;
    color: white;
    padding: 12px 22px;
    border: none;
    border-radius: 8px;
    font-weight: 700;
    cursor: pointer;
    font-family: 'Cairo', sans-serif;
    font-size: 14px;
    box-shadow: 0 4px 12px rgba(0,168,50,0.3);
    z-index: 1000;
}}
.print-btn:hover {{ background: #008a28; }}
.footer {{ margin-top: 50px; padding-top: 20px; border-top: 1px solid #e5e7eb; text-align: center; color: #9ca3af; font-size: 12px; }}
@media print {{ body {{ background: white; margin: 0; }} .print-btn {{ display: none; }} .finding {{ break-inside: avoid; page-break-inside: avoid; }} }}
</style>
</head>
<body>
<button class="print-btn" onclick="window.print()">🖨️ طباعة / حفظ PDF</button>
<h1>🛡️ تقرير فحص أمني — 10blackfingers</h1>

<div class="overall {overall_class}">{_esc(overall_msg)}</div>

<div class="summary">
    <p><strong>الهدف:</strong> {target_safe}</p>
    <p><strong>رقم الفحص:</strong> {_esc(scan.get('id', 'N/A'))}</p>
    <p><strong>تاريخ البدء:</strong> {_esc(scan.get('started_at', 'N/A'))}</p>
    <p><strong>تاريخ الانتهاء:</strong> {_esc(scan.get('completed_at', 'N/A'))}</p>
    <p><strong>المدة:</strong> {_esc(scan.get('elapsed_seconds', 'N/A'))} ثانية</p>
    <div class="stats">
        <div class="stat total"><div class="num">{total}</div><div class="lbl">إجمالي</div></div>
        <div class="stat critical"><div class="num">{critical}</div><div class="lbl">حرجة</div></div>
        <div class="stat high"><div class="num">{high}</div><div class="lbl">عالية</div></div>
        <div class="stat medium"><div class="num">{medium}</div><div class="lbl">متوسطة</div></div>
        <div class="stat low"><div class="num">{low}</div><div class="lbl">منخفضة</div></div>
    </div>
</div>

{findings_html}

<div class="footer">
    تم توليد هذا التقرير بواسطة منصة 10blackfingers في {datetime.utcnow().isoformat()} UTC
</div>
</body>
</html>"""
    return html_out


def build_json_report(scan, results):
    return {
        "scan": scan,
        "results": results,
        "generated_at": datetime.utcnow().isoformat(),
        "generator": "10blackfingers v2.1.0",
    }


@app.get("/api/scan/{scan_id}/report/txt")
async def download_report_txt(scan_id: str, session_id: str = Query(...)):
    scan, results = await fetch_report_data(scan_id, session_id)
    content = build_txt_report(scan, results)
    return PlainTextResponse(
        content,
        headers=secure_report_headers(scan_id, "txt", attachment=True),
    )


@app.get("/api/scan/{scan_id}/report/html", response_class=HTMLResponse)
async def download_report_html(scan_id: str, session_id: str = Query(...)):
    scan, results = await fetch_report_data(scan_id, session_id)
    return HTMLResponse(
        build_html_report(scan, results),
        headers=secure_report_headers(scan_id, "html", attachment=False),
    )


@app.get("/api/scan/{scan_id}/report/json")
async def download_report_json(scan_id: str, session_id: str = Query(...)):
    scan, results = await fetch_report_data(scan_id, session_id)
    return JSONResponse(
        content=build_json_report(scan, results),
        headers=secure_report_headers(scan_id, "json", attachment=True),
    )


@app.get("/api/scan/{scan_id}/report/xlsx")
async def download_report_xlsx(scan_id: str, session_id: str = Query(...)):
    scan, results = await fetch_report_data(scan_id, session_id)

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError as e:
        raise HTTPException(status_code=500, detail=f"Excel library not available: {str(e)}")

    wb = Workbook()

    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Cairo", bold=True, color="FFFFFF", size=11)

    severity_fills = {
        "critical": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
        "high": PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid"),
        "medium": PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid"),
        "low": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
        "info": PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid"),
    }

    ws = wb.active
    ws.title = "الملخص"
    ws.sheet_view.rightToLeft = True

    ws["A1"] = "10blackfingers — تقرير فحص أمني"
    ws["A1"].font = Font(name="Cairo", bold=True, size=16, color="00A832")
    ws.merge_cells("A1:B1")

    summary_data = [
        ("الهدف", scan.get("target_url", "N/A")),
        ("رقم الفحص", scan.get("id", "N/A")),
        ("تاريخ البدء", scan.get("started_at", "N/A")),
        ("تاريخ الانتهاء", scan.get("completed_at", "N/A")),
        ("المدة (ثانية)", scan.get("elapsed_seconds", "N/A")),
        ("الحالة", scan.get("status", "N/A")),
        ("إجمالي النتائج", scan.get("total_findings", 0)),
        ("حرجة", scan.get("critical_count", 0)),
        ("عالية", scan.get("high_count", 0)),
        ("متوسطة", scan.get("medium_count", 0)),
        ("منخفضة", scan.get("low_count", 0)),
    ]

    row = 3
    for key, value in summary_data:
        ws.cell(row=row, column=1, value=key).font = Font(name="Cairo", bold=True)
        ws.cell(row=row, column=2, value=str(value)).font = Font(name="Cairo")
        ws.cell(row=row, column=1).fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        ws.cell(row=row, column=1).border = thin_border
        ws.cell(row=row, column=2).border = thin_border
        ws.cell(row=row, column=1).alignment = Alignment(horizontal="right", vertical="center")
        ws.cell(row=row, column=2).alignment = Alignment(horizontal="left", vertical="center")
        row += 1

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 60

    ws2 = wb.create_sheet("الثغرات")
    ws2.sheet_view.rightToLeft = True

    headers = ["#", "الخطورة", "العنوان", "الأداة", "المرحلة", "CVSS", "CVE", "الوصف", "الأدلة", "الإصلاح"]
    for col_idx, header in enumerate(headers, 1):
        cell = ws2.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    sev_order = {"critical": 0, "high": 1, "medium": 2, "low": 3, "info": 4}
    sorted_results = sorted(results, key=lambda r: sev_order.get(r.get("severity", "info"), 5))

    for idx, r in enumerate(sorted_results, 1):
        sev = r.get("severity", "info")
        evidence_text = ""
        ev = r.get("evidence")
        if isinstance(ev, dict):
            parts = []
            for k, v in ev.items():
                if k in ("extracted_pii", "aggregated", "pii_data", "dumped_files_preview", "extracted_files"):
                    parts.append(f"[{k}]: {_fmt_val(v, 2000)}")
                else:
                    parts.append(f"{k}: {_fmt_val(v, 500)}")
            evidence_text = "\n".join(parts)
        elif isinstance(ev, str):
            evidence_text = ev[:5000]

        row_data = [
            idx,
            sev,
            r.get("title", ""),
            r.get("tool_name", ""),
            r.get("phase", ""),
            r.get("cvss_score", "") or "",
            r.get("cve_id", "") or "",
            (r.get("description", "") or "")[:2000],
            evidence_text,
            (r.get("remediation", "") or "")[:1000],
        ]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws2.cell(row=idx + 1, column=col_idx, value=value)
            cell.font = Font(name="Cairo", size=9)
            cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
            cell.border = thin_border
        if sev in severity_fills:
            ws2.cell(row=idx + 1, column=2).fill = severity_fills[sev]

    col_widths = [6, 12, 40, 15, 15, 10, 20, 50, 60, 40]
    for i, w in enumerate(col_widths, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "A2"

    ws3 = wb.create_sheet("البيانات الحساسة")
    ws3.sheet_view.rightToLeft = True

    pii_headers = ["#", "النوع", "القيمة", "المصدر", "الثغرة المرتبطة"]
    for col_idx, header in enumerate(pii_headers, 1):
        cell = ws3.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    pii_row = 2
    pii_counter = 1
    for r in results:
        ev = r.get("evidence")
        if not isinstance(ev, dict):
            continue
        finding_title = r.get("title", "")
        for key in ("extracted_pii", "aggregated"):
            if key in ev and isinstance(ev[key], dict):
                for pk, pv in ev[key].items():
                    if pk == "source_url":
                        continue
                    items = pv if isinstance(pv, list) else [pv]
                    for item in items:
                        ws3.cell(row=pii_row, column=1, value=pii_counter).font = Font(name="Cairo", size=9)
                        ws3.cell(row=pii_row, column=2, value=str(pk)).font = Font(name="Cairo", size=9)
                        ws3.cell(row=pii_row, column=3, value=str(item)).font = Font(name="Courier New", size=9)
                        ws3.cell(row=pii_row, column=4, value="").font = Font(name="Cairo", size=9)
                        ws3.cell(row=pii_row, column=5, value=finding_title).font = Font(name="Cairo", size=9)
                        for c in range(1, 6):
                            ws3.cell(row=pii_row, column=c).border = thin_border
                            ws3.cell(row=pii_row, column=c).alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
                        pii_row += 1
                        pii_counter += 1
        if "pii_data" in ev and isinstance(ev["pii_data"], list):
            for item in ev["pii_data"]:
                if not isinstance(item, dict):
                    continue
                source = item.get("source_url", "")
                for pk, pv in item.items():
                    if pk == "source_url":
                        continue
                    items = pv if isinstance(pv, list) else [pv]
                    for sub in items:
                        ws3.cell(row=pii_row, column=1, value=pii_counter).font = Font(name="Cairo", size=9)
                        ws3.cell(row=pii_row, column=2, value=str(pk)).font = Font(name="Cairo", size=9)
                        ws3.cell(row=pii_row, column=3, value=str(sub)).font = Font(name="Courier New", size=9)
                        ws3.cell(row=pii_row, column=4, value=str(source)).font = Font(name="Cairo", size=9)
                        ws3.cell(row=pii_row, column=5, value=finding_title).font = Font(name="Cairo", size=9)
                        for c in range(1, 6):
                            ws3.cell(row=pii_row, column=c).border = thin_border
                            ws3.cell(row=pii_row, column=c).alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
                        pii_row += 1
                        pii_counter += 1

    if pii_row == 2:
        ws3.cell(row=2, column=1, value="لا توجد بيانات حساسة مستخرجة").font = Font(name="Cairo", italic=True)

    pii_col_widths = [6, 20, 50, 40, 40]
    for i, w in enumerate(pii_col_widths, 1):
        ws3.column_dimensions[get_column_letter(i)].width = w
    ws3.freeze_panes = "A2"

    ws4 = wb.create_sheet("قاعدة البيانات")
    ws4.sheet_view.rightToLeft = True

    db_headers = ["#", "الملف", "المحتوى (مقتطع)"]
    for col_idx, header in enumerate(db_headers, 1):
        cell = ws4.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    db_row = 2
    db_counter = 1
    for r in results:
        ev = r.get("evidence")
        if not isinstance(ev, dict):
            continue
        dumped = ev.get("dumped_files_preview")
        if isinstance(dumped, dict):
            for filename, content in dumped.items():
                ws4.cell(row=db_row, column=1, value=db_counter).font = Font(name="Cairo", size=9)
                ws4.cell(row=db_row, column=2, value=str(filename)).font = Font(name="Cairo", size=9)
                ws4.cell(row=db_row, column=3, value=str(content)[:30000]).font = Font(name="Courier New", size=8)
                for c in range(1, 4):
                    ws4.cell(row=db_row, column=c).border = thin_border
                    ws4.cell(row=db_row, column=c).alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
                db_row += 1
                db_counter += 1

    if db_row == 2:
        ws4.cell(row=2, column=1, value="لا توجد بيانات قاعدة بيانات مسحوبة").font = Font(name="Cairo", italic=True)

    ws4.column_dimensions["A"].width = 6
    ws4.column_dimensions["B"].width = 40
    ws4.column_dimensions["C"].width = 100
    ws4.freeze_panes = "A2"

    ws5 = wb.create_sheet("الملفات المسحوبة")
    ws5.sheet_view.rightToLeft = True

    files_headers = ["#", "الملف", "الرابط", "المحتوى (مقتطع)"]
    for col_idx, header in enumerate(files_headers, 1):
        cell = ws5.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    files_row = 2
    files_counter = 1
    for r in results:
        ev = r.get("evidence")
        if not isinstance(ev, dict):
            continue
        files_list = ev.get("extracted_files")
        if isinstance(files_list, list):
            for f in files_list:
                if not isinstance(f, dict):
                    continue
                ws5.cell(row=files_row, column=1, value=files_counter).font = Font(name="Cairo", size=9)
                ws5.cell(row=files_row, column=2, value=str(f.get("file", ""))).font = Font(name="Cairo", size=9)
                ws5.cell(row=files_row, column=3, value=str(f.get("url", ""))).font = Font(name="Courier New", size=9)
                ws5.cell(row=files_row, column=4, value=str(f.get("content", ""))[:30000]).font = Font(name="Courier New", size=8)
                for c in range(1, 5):
                    ws5.cell(row=files_row, column=c).border = thin_border
                    ws5.cell(row=files_row, column=c).alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
                files_row += 1
                files_counter += 1

    if files_row == 2:
        ws5.cell(row=2, column=1, value="لا توجد ملفات مسحوبة").font = Font(name="Cairo", italic=True)

    ws5.column_dimensions["A"].width = 6
    ws5.column_dimensions["B"].width = 30
    ws5.column_dimensions["C"].width = 60
    ws5.column_dimensions["D"].width = 100
    ws5.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    xlsx_bytes = buffer.getvalue()
    buffer.close()

    return Response(
        content=xlsx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=secure_report_headers(scan_id, "xlsx", attachment=True),
    )


@app.get("/api/scan/{scan_id}/report/pdf")
async def download_report_pdf(scan_id: str, session_id: str = Query(...)):
    scan, results = await fetch_report_data(scan_id, session_id)

    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        import arabic_reshaper
        from bidi.algorithm import get_display
    except ImportError as e:
        raise HTTPException(status_code=500, detail=f"PDF library not available: {str(e)}")

    font_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fonts", "Cairo-Regular.ttf")
    if not os.path.exists(font_path):
        raise HTTPException(
            status_code=500,
            detail="ملف الخط العربي (Cairo-Regular.ttf) غير موجود - تحقق من إعدادات includeFiles في vercel.json",
        )

    try:
        pdfmetrics.registerFont(TTFont("Cairo", font_path))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"تعذر تحميل الخط العربي: {str(e)}")

    arabic_font_name = "Cairo"

    def ar(text):
        if not text:
            return ""
        try:
            reshaped = arabic_reshaper.reshape(str(text))
            bidi_text = get_display(reshaped)
            return (bidi_text
                    .replace("&", "&amp;")
                    .replace("<", "&lt;")
                    .replace(">", "&gt;"))
        except Exception:
            return (str(text)
                    .replace("&", "&amp;")
                    .replace("<", "&lt;")
                    .replace(">", "&gt;"))

    def en(text):
        return (str(text or "")
                .replace("&", "&amp;")
                .replace("<", "&lt;")
                .replace(">", "&gt;"))

    def smart_text(text, max_len=None):
        if not text:
            return ""
        s = str(text)
        if max_len and len(s) > max_len:
            s = s[:max_len]
        if any("\u0600" <= c <= "\u06ff" for c in s):
            return ar(s)
        return en(s)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4,
        leftMargin=15 * mm, rightMargin=15 * mm,
        topMargin=15 * mm, bottomMargin=15 * mm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'T', parent=styles['Title'], fontSize=22,
        textColor=colors.HexColor('#00a832'), fontName=arabic_font_name,
        alignment=1,
    )
    h2_style = ParagraphStyle(
        'H2', parent=styles['Heading2'], fontSize=14,
        spaceAfter=10, fontName=arabic_font_name,
        textColor=colors.HexColor('#111111'),
    )
    h3_style = ParagraphStyle(
        'H3', parent=styles['Heading3'], fontSize=11,
        spaceAfter=6, fontName=arabic_font_name,
    )
    body_style = ParagraphStyle(
        'B', parent=styles['Normal'], fontSize=9, leading=13,
        fontName=arabic_font_name,
    )
    body_rtl = ParagraphStyle(
        'BR', parent=body_style, alignment=2,
    )

    story = []
    story.append(Paragraph(ar("تقرير فحص أمني - 10blackfingers"), title_style))
    story.append(Spacer(1, 8))

    total = scan.get('total_findings', 0)
    critical = scan.get('critical_count', 0)
    high = scan.get('high_count', 0)
    medium = scan.get('medium_count', 0)
    low = scan.get('low_count', 0)

    if critical > 0:
        overall = "خطر - يوجد ثغرات حرجة تحتاج إصلاح فوري"
    elif high > 0:
        overall = "يحتاج انتباه - يوجد ثغرات عالية"
    elif medium > 0:
        overall = "متوسط - يوجد ثغرات متوسطة الخطورة"
    else:
        overall = "آمن - لا يوجد ثغرات خطيرة"

    info_text = (
        f"<b>{ar('الهدف:')}</b> {en(scan.get('target_url', 'N/A'))}<br/>"
        f"<b>{ar('رقم الفحص:')}</b> {en(scan.get('id', 'N/A'))}<br/>"
        f"<b>{ar('المدة:')}</b> {en(scan.get('elapsed_seconds', 'N/A'))} {ar('ثانية')}<br/>"
        f"<b>{ar('التقييم العام:')}</b> {ar(overall)}<br/>"
        f"<b>{ar('إجمالي النتائج:')}</b> {total}<br/>"
        f"<b>{ar('حرجة:')}</b> {critical} | "
        f"<b>{ar('عالية:')}</b> {high} | "
        f"<b>{ar('متوسطة:')}</b> {medium} | "
        f"<b>{ar('منخفضة:')}</b> {low}"
    )
    story.append(Paragraph(info_text, body_rtl))
    story.append(Spacer(1, 20))

    by_severity = {"critical": [], "high": [], "medium": [], "low": [], "info": []}
    for r in results:
        by_severity.setdefault(r.get("severity", "info"), []).append(r)

    severity_ar = {
        "critical": "ثغرات حرجة",
        "high": "ثغرات عالية",
        "medium": "ثغرات متوسطة",
        "low": "ثغرات منخفضة",
        "info": "معلومات",
    }

    for severity in ["critical", "high", "medium", "low", "info"]:
        items = by_severity.get(severity, [])
        if not items:
            continue

        story.append(Paragraph(ar(f"{severity_ar.get(severity, severity)} ({len(items)})"), h2_style))
        story.append(Spacer(1, 6))

        for i, r in enumerate(items, 1):
            title = r.get('title', 'Untitled')
            story.append(Paragraph(f"[{i}] {smart_text(title, max_len=300)}", h3_style))

            meta = f"<b>Tool:</b> {en(r.get('tool_name', 'N/A'))} | <b>Phase:</b> {en(r.get('phase', 'N/A'))}"
            if r.get('cvss_score'):
                meta += f" | <b>CVSS:</b> {en(r['cvss_score'])}"
            if r.get('cve_id'):
                meta += f" | <b>CVE:</b> {en(r['cve_id'])}"
            story.append(Paragraph(meta, body_style))

            if r.get('description'):
                desc = smart_text(r['description'], max_len=1500)
                story.append(Paragraph(f"<b>{ar('الوصف:')}</b> {desc}", body_rtl))

            ev = r.get('evidence')
            if isinstance(ev, dict):
                ev_text = f"<b>{ar('الأدلة:')}</b><br/>"
                count = 0
                for k, v in ev.items():
                    if count >= 8:
                        break
                    if k in ("extracted_pii", "aggregated") and isinstance(v, dict):
                        ev_text += f"&nbsp;&nbsp;- <b>{en(k)}:</b><br/>"
                        sub_count = 0
                        for pk, pv in v.items():
                            if sub_count >= 5:
                                break
                            if isinstance(pv, list):
                                vals = ", ".join(str(x) for x in pv[:3])
                            else:
                                vals = str(pv)[:200]
                            ev_text += f"&nbsp;&nbsp;&nbsp;&nbsp;{en(pk)}: {en(vals)[:300]}<br/>"
                            sub_count += 1
                    elif k == "dumped_files_preview" and isinstance(v, dict):
                        ev_text += f"&nbsp;&nbsp;- <b>DB Files ({len(v)}):</b> {en(', '.join(list(v.keys())[:5]))}<br/>"
                    elif k == "extracted_files" and isinstance(v, list):
                        names = ", ".join(f.get('file', '') for f in v[:5] if isinstance(f, dict))
                        ev_text += f"&nbsp;&nbsp;- <b>Files ({len(v)}):</b> {en(names)}<br/>"
                    else:
                        ev_text += f"&nbsp;&nbsp;- {en(k)}: {en(str(v)[:300])}<br/>"
                    count += 1
                story.append(Paragraph(ev_text, body_style))
            elif isinstance(ev, str) and ev:
                story.append(Paragraph(f"<b>{ar('الأدلة:')}</b> {en(ev[:500])}", body_rtl))

            if r.get('remediation'):
                rem = smart_text(r['remediation'], max_len=800)
                story.append(Paragraph(f"<b>{ar('الإصلاح:')}</b> {rem}", body_rtl))

            story.append(Spacer(1, 10))

        story.append(PageBreak())

    try:
        doc.build(story)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"PDF build error: {str(e)}")

    pdf_bytes = buffer.getvalue()
    buffer.close()

    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers=secure_report_headers(scan_id, "pdf", attachment=True),
    )